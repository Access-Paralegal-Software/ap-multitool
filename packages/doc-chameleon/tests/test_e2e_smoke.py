"""End-to-end smoke tests — full convert pipeline for all fixture/jurisdiction combos.

These tests exercise the complete path from source fixture through transform,
save, and report write, then verify the outputs are well-formed and loadable.
They are intentionally broad, not exhaustive — unit tests cover the details.
"""
from __future__ import annotations

from pathlib import Path

import docx as python_docx

from doc_chameleon.engine.export import save
from doc_chameleon.engine.ingest import load
from doc_chameleon.engine.report import write_report
from doc_chameleon.rules.ca.transformer import transform as transform_ca
from doc_chameleon.rules.ca.transformer_2111 import transform_2111 as transform_ca_2111
from doc_chameleon.rules.ca.validator import validate as validate_ca
from doc_chameleon.rules.ca.validator import validate_2111 as validate_ca_2111
from doc_chameleon.rules.ca.validator import validate_source_assumptions as validate_ca_source_assumptions
from doc_chameleon.rules.tx.transformer import transform as transform_tx
from doc_chameleon.rules.tx.validator import validate as validate_tx
from doc_chameleon.rules.tx.validator import validate_source_assumptions as validate_tx_source_assumptions

from tests.fixtures import (
    make_cover_page,
    make_declaration,
    make_hostile,
    make_motion,
    make_notice,
    make_tx_motion,
    make_tx_notice,
)


def _run_ca(tmp_path: Path, fixture_name: str, make_fn) -> tuple[Path, Path, list[str]]:
    """Run the full CA pipeline on a fixture and return (output_path, report_path, warnings)."""
    src = tmp_path / f"{fixture_name}.docx"
    out = tmp_path / f"{fixture_name}_ca.docx"
    make_fn(src)

    doc, record = load(src)
    warnings: list[str] = []
    warnings.extend(validate_ca_source_assumptions(doc, record))
    doc = transform_ca(doc)
    doc = transform_ca_2111(doc)
    warnings.extend(validate_ca(doc, record))
    warnings.extend(validate_ca_2111(doc, record))
    warnings = list(dict.fromkeys(warnings))

    save(doc, out)
    report = write_report(out, src, "ca", warnings)
    return out, report, warnings


def _run_tx(tmp_path: Path, fixture_name: str, make_fn) -> tuple[Path, Path, list[str]]:
    """Run the full TX pipeline on a fixture and return (output_path, report_path, warnings)."""
    src = tmp_path / f"{fixture_name}.docx"
    out = tmp_path / f"{fixture_name}_tx.docx"
    make_fn(src)

    doc, record = load(src)
    warnings: list[str] = []
    warnings.extend(validate_tx_source_assumptions(doc, record))
    doc = transform_tx(doc)
    warnings.extend(validate_tx(doc, record))
    warnings = list(dict.fromkeys(warnings))

    save(doc, out)
    report = write_report(out, src, "tx", warnings)
    return out, report, warnings


def _assert_output_well_formed(out: Path, report: Path) -> None:
    assert out.exists(), f"Output .docx not found: {out}"
    assert report.exists(), f"Report .txt not found: {report}"
    reloaded = python_docx.Document(str(out))
    assert reloaded.sections, "Output .docx has no sections"
    content = report.read_text(encoding="utf-8")
    assert "doc-chameleon Conversion Report" in content
    assert "does not provide legal advice" in content


# ── CA smoke tests ────────────────────────────────────────────────────────────

def test_smoke_ca_motion(tmp_path: Path) -> None:
    out, report, warnings = _run_ca(tmp_path, "motion", make_motion)
    _assert_output_well_formed(out, report)
    assert warnings == [], f"Clean motion should produce no CA warnings: {warnings}"


def test_smoke_ca_declaration(tmp_path: Path) -> None:
    out, report, warnings = _run_ca(tmp_path, "declaration", make_declaration)
    _assert_output_well_formed(out, report)
    assert warnings == [], f"Clean declaration should produce no CA warnings: {warnings}"


def test_smoke_ca_notice(tmp_path: Path) -> None:
    out, report, warnings = _run_ca(tmp_path, "notice", make_notice)
    _assert_output_well_formed(out, report)
    assert warnings == [], f"Clean notice should produce no CA warnings: {warnings}"


def test_smoke_ca_cover_page(tmp_path: Path) -> None:
    out, report, warnings = _run_ca(tmp_path, "cover_page", make_cover_page)
    _assert_output_well_formed(out, report)
    assert warnings == [], f"Clean cover page should produce no CA warnings: {warnings}"


def test_smoke_ca_hostile_does_not_crash(tmp_path: Path) -> None:
    """Hostile layout must produce output and warnings, but must not crash."""
    out, report, warnings = _run_ca(tmp_path, "hostile", make_hostile)
    _assert_output_well_formed(out, report)
    assert warnings, "Hostile layout must produce at least one warning"
    assert any("Tables" in w for w in warnings)
    report_content = report.read_text(encoding="utf-8")
    assert "Warnings" in report_content


def test_smoke_ca_report_has_correct_metadata(tmp_path: Path) -> None:
    out, report, _ = _run_ca(tmp_path, "motion", make_motion)
    content = report.read_text(encoding="utf-8")
    assert "motion.docx" in content
    assert "motion_ca.docx" in content
    assert "ca" in content


# ── TX smoke tests ────────────────────────────────────────────────────────────

def test_smoke_tx_motion(tmp_path: Path) -> None:
    out, report, warnings = _run_tx(tmp_path, "tx_motion", make_tx_motion)
    _assert_output_well_formed(out, report)
    # TX always includes the statewide-only disclaimer — that is expected
    non_disclaimer = [w for w in warnings if "statewide baseline" not in w]
    assert non_disclaimer == [], f"TX motion should have no warnings beyond the disclaimer: {non_disclaimer}"


def test_smoke_tx_notice(tmp_path: Path) -> None:
    out, report, warnings = _run_tx(tmp_path, "tx_notice", make_tx_notice)
    _assert_output_well_formed(out, report)
    non_disclaimer = [w for w in warnings if "statewide baseline" not in w]
    assert non_disclaimer == [], f"TX notice should have no warnings beyond the disclaimer: {non_disclaimer}"


def test_smoke_tx_report_includes_statewide_disclaimer(tmp_path: Path) -> None:
    _, report, _ = _run_tx(tmp_path, "tx_motion", make_tx_motion)
    content = report.read_text(encoding="utf-8")
    assert "statewide baseline" in content


def test_smoke_tx_report_has_correct_metadata(tmp_path: Path) -> None:
    out, report, _ = _run_tx(tmp_path, "tx_motion", make_tx_motion)
    content = report.read_text(encoding="utf-8")
    assert "tx_motion.docx" in content
    assert "tx_motion_tx.docx" in content
    assert "tx" in content


# ── cross-jurisdiction guard ──────────────────────────────────────────────────

def test_smoke_ca_output_has_no_tx_marker(tmp_path: Path) -> None:
    """A CA-converted document must not contain the TX transform marker."""
    from doc_chameleon.rules.tx.transformer import TX_MARKER
    out, _, _ = _run_ca(tmp_path, "motion", make_motion)
    reloaded = python_docx.Document(str(out))
    assert TX_MARKER not in reloaded.sections[0]._sectPr.xml


def test_smoke_tx_output_has_no_ca_line_numbers(tmp_path: Path) -> None:
    """A TX-converted document must not contain CA line-number headers."""
    out, _, _ = _run_tx(tmp_path, "tx_motion", make_tx_motion)
    reloaded = python_docx.Document(str(out))
    for section in reloaded.sections:
        assert "doc-chameleon-ca-line-numbers" not in section.header._element.xml
