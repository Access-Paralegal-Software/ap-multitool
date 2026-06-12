"""Synthetic .docx fixture factories for doc-chameleon tests."""
from __future__ import annotations

from pathlib import Path

import docx
from docx.shared import Pt


def make_motion(path: Path) -> None:
    """Clean motion for summary judgment — no hostile layout elements."""
    doc = docx.Document()
    doc.add_paragraph("MOTION FOR SUMMARY JUDGMENT")
    doc.add_paragraph("")
    doc.add_paragraph(
        "Plaintiff Jane Smith ('Plaintiff') respectfully moves this Court pursuant to "
        "Code of Civil Procedure section 437c for summary judgment on all causes of "
        "action asserted in the operative complaint."
    )
    doc.add_paragraph("INTRODUCTION")
    doc.add_paragraph(
        "The material facts are undisputed. Defendant breached the written contract "
        "dated January 1, 2025, by failing to deliver the goods described therein."
    )
    doc.add_paragraph("STATEMENT OF UNDISPUTED MATERIAL FACTS")
    for i in range(1, 8):
        doc.add_paragraph(
            f"{i}. Undisputed fact number {i} is established by the evidence submitted herewith."
        )
    doc.add_paragraph("ARGUMENT")
    doc.add_paragraph(
        "Under governing California law, summary judgment is proper where no triable "
        "issue of material fact exists and the moving party is entitled to judgment "
        "as a matter of law. (Code Civ. Proc., § 437c, subd. (c).)"
    )
    doc.add_paragraph("CONCLUSION")
    doc.add_paragraph(
        "For the foregoing reasons, Plaintiff respectfully requests that the Court "
        "grant summary judgment in her favor on all causes of action."
    )
    doc.save(str(path))


def make_declaration(path: Path) -> None:
    """Declaration with numbered paragraphs — clean layout, no hostile elements."""
    doc = docx.Document()
    doc.add_paragraph(
        "DECLARATION OF JANE SMITH IN SUPPORT OF MOTION FOR SUMMARY JUDGMENT"
    )
    doc.add_paragraph("I, Jane Smith, declare as follows:")
    for i in range(1, 12):
        doc.add_paragraph(f"{i}. Paragraph {i} of the declaration sets forth relevant facts.")
    doc.add_paragraph(
        "I declare under penalty of perjury under the laws of the State of California "
        "that the foregoing is true and correct."
    )
    doc.add_paragraph("Executed on June 12, 2026, at Los Angeles, California.")
    doc.add_paragraph("")
    doc.add_paragraph("Jane Smith")
    doc.save(str(path))


def make_notice(path: Path) -> None:
    """Short notice of motion — minimal content, clean layout."""
    doc = docx.Document()
    doc.add_paragraph("NOTICE OF MOTION AND MOTION FOR SUMMARY JUDGMENT")
    doc.add_paragraph("TO ALL PARTIES AND THEIR ATTORNEYS OF RECORD:")
    doc.add_paragraph(
        "PLEASE TAKE NOTICE that on September 15, 2026, at 9:00 a.m., or as soon "
        "thereafter as the matter may be heard, in Department 32 of the above-entitled "
        "court, Plaintiff Jane Smith will move this Court for summary judgment."
    )
    doc.add_paragraph(
        "This motion will be based on this notice, the attached memorandum of points "
        "and authorities, the declaration of Jane Smith, and all pleadings and records "
        "on file herein."
    )
    doc.save(str(path))


def make_cover_page(path: Path) -> None:
    """Pleading cover page with court title and case caption — clean layout."""
    doc = docx.Document()
    doc.add_paragraph("IN THE SUPERIOR COURT OF THE STATE OF CALIFORNIA")
    doc.add_paragraph("FOR THE COUNTY OF LOS ANGELES")
    doc.add_paragraph("")
    doc.add_paragraph("JANE SMITH,")
    doc.add_paragraph("    Plaintiff,")
    doc.add_paragraph("")
    doc.add_paragraph("    v.")
    doc.add_paragraph("")
    doc.add_paragraph("JOHN DOE,")
    doc.add_paragraph("    Defendant.")
    doc.add_paragraph("")
    doc.add_paragraph("MOTION FOR SUMMARY JUDGMENT")
    doc.save(str(path))


def make_hostile(path: Path) -> None:
    """Hostile layout: explicit paragraph spacing and an embedded table."""
    doc = docx.Document()
    hostile_para = doc.add_paragraph("Document with non-standard layout spacing.")
    hostile_para.paragraph_format.space_after = Pt(12)
    hostile_para.paragraph_format.space_before = Pt(6)
    doc.add_paragraph("This paragraph has default Word spacing applied.")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Column A"
    table.cell(0, 1).text = "Column B"
    table.cell(1, 0).text = "Value 1"
    table.cell(1, 1).text = "Value 2"
    doc.add_paragraph("Post-table paragraph with trailing content.")
    doc.save(str(path))
