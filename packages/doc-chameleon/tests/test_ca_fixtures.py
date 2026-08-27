"""CA transformer tests against synthetic fixture document classes."""
from __future__ import annotations

from pathlib import Path
from zipfile import ZipFile

from docx.enum.text import WD_LINE_SPACING
from docx.shared import Inches, Pt

from doc_chameleon.engine.export import save
from doc_chameleon.engine.ingest import load
from doc_chameleon.engine.report import write_report
from doc_chameleon.rules.ca.transformer import PLEADING_LINE_COUNT, transform
from doc_chameleon.rules.ca.validator import validate, validate_source_assumptions

from tests.fixtures import make_declaration, make_hostile, make_motion, make_notice


def _has_ca_header(path: Path) -> bool:
    with ZipFile(path) as zf:
        header_names = [n for n in zf.namelist() if n.startswith("word/header")]
        if not header_names:
            return False
        return any(
            "doc-chameleon-ca-line-numbers" in zf.read(n).decode("utf-8")
            for n in header_names
        )


def _assert_ca_geometry(doc) -> None:
    for section in doc.sections:
        assert section.page_width == Inches(8.5)
        assert section.page_height == Inches(11)
        assert section.left_margin == Inches(1.25)
        assert section.right_margin == Inches(1)
        assert section.top_margin == Inches(1)
        assert section.bottom_margin == Inches(1)


def _assert_ca_body_flow(doc) -> None:
    for paragraph in doc.paragraphs:
        assert paragraph.style.name == "Doc Chameleon CA Body"
        assert paragraph.paragraph_format.line_spacing_rule == WD_LINE_SPACING.EXACTLY
        assert paragraph.paragraph_format.line_spacing == Pt(24)
        assert paragraph.paragraph_format.space_before == Pt(0)
        assert paragraph.paragraph_format.space_after == Pt(0)


# ── motion fixture ────────────────────────────────────────────────────────────

def test_motion_transforms_cleanly(tmp_path: Path) -> None:
    src = tmp_path / "motion.docx"
    out = tmp_path / "motion_ca.docx"
    make_motion(src)

    doc, record = load(src)
    source_warnings = validate_source_assumptions(doc, record)
    assert source_warnings == [], f"Clean motion should have no source warnings: {source_warnings}"

    transform(doc)
    save(doc, out)

    assert _has_ca_header(out)
    reloaded, _ = load(out)
    _assert_ca_geometry(reloaded)
    _assert_ca_body_flow(reloaded)


def test_motion_produces_no_post_transform_geometry_warnings(tmp_path: Path) -> None:
    src = tmp_path / "motion.docx"
    out = tmp_path / "motion_ca.docx"
    make_motion(src)

    doc, record = load(src)
    transform(doc)
    save(doc, out)

    reloaded, reloaded_record = load(out)
    warnings = validate(reloaded, reloaded_record)
    geometry_warnings = [w for w in warnings if "margin" in w or "page width" in w or "page height" in w]
    assert geometry_warnings == [], f"No geometry warnings expected after transform: {geometry_warnings}"


# ── declaration fixture ───────────────────────────────────────────────────────

def test_declaration_transforms_cleanly(tmp_path: Path) -> None:
    src = tmp_path / "declaration.docx"
    out = tmp_path / "declaration_ca.docx"
    make_declaration(src)

    doc, record = load(src)
    source_warnings = validate_source_assumptions(doc, record)
    assert source_warnings == [], f"Clean declaration should have no source warnings: {source_warnings}"

    transform(doc)
    save(doc, out)

    assert _has_ca_header(out)
    reloaded, _ = load(out)
    _assert_ca_geometry(reloaded)
    _assert_ca_body_flow(reloaded)


def test_declaration_preserves_paragraph_count(tmp_path: Path) -> None:
    src = tmp_path / "declaration.docx"
    out = tmp_path / "declaration_ca.docx"
    make_declaration(src)

    doc, record = load(src)
    original_count = len(record.paragraphs)
    transform(doc)
    save(doc, out)

    reloaded, reloaded_record = load(out)
    assert len(reloaded_record.paragraphs) == original_count


# ── notice fixture ────────────────────────────────────────────────────────────

def test_notice_transforms_cleanly(tmp_path: Path) -> None:
    src = tmp_path / "notice.docx"
    out = tmp_path / "notice_ca.docx"
    make_notice(src)

    doc, record = load(src)
    source_warnings = validate_source_assumptions(doc, record)
    assert source_warnings == [], f"Clean notice should have no source warnings: {source_warnings}"

    transform(doc)
    save(doc, out)

    assert _has_ca_header(out)
    reloaded, _ = load(out)
    _assert_ca_geometry(reloaded)


# ── hostile layout fixture ────────────────────────────────────────────────────

def test_hostile_layout_triggers_source_warnings(tmp_path: Path) -> None:
    src = tmp_path / "hostile.docx"
    make_hostile(src)

    doc, record = load(src)
    warnings = validate_source_assumptions(doc, record)

    warning_text = " ".join(warnings)
    assert "Tables detected" in warning_text
    assert "paragraph spacing" in warning_text


def test_hostile_layout_still_transforms_and_saves(tmp_path: Path) -> None:
    """Transformer must not crash on hostile input — it warns, does not refuse."""
    src = tmp_path / "hostile.docx"
    out = tmp_path / "hostile_ca.docx"
    make_hostile(src)

    doc, record = load(src)
    transform(doc)
    save(doc, out)

    assert out.exists()
    assert _has_ca_header(out)


# ── conversion report ─────────────────────────────────────────────────────────

def test_conversion_report_written_alongside_output(tmp_path: Path) -> None:
    src = tmp_path / "motion.docx"
    out = tmp_path / "motion_ca.docx"
    make_motion(src)

    doc, record = load(src)
    transform(doc)
    save(doc, out)
    report_path = write_report(out, src, "ca", [])

    assert report_path == tmp_path / "motion_ca_report.txt"
    assert report_path.exists()
    content = report_path.read_text(encoding="utf-8")
    assert "doc-chameleon Conversion Report" in content
    assert "motion.docx" in content
    assert "motion_ca.docx" in content
    assert "ca" in content


def test_conversion_report_includes_warnings_when_present(tmp_path: Path) -> None:
    src = tmp_path / "hostile.docx"
    out = tmp_path / "hostile_ca.docx"
    make_hostile(src)

    doc, record = load(src)
    source_warnings = validate_source_assumptions(doc, record)
    transform(doc)
    save(doc, out)
    warnings = validate(doc, record)
    all_warnings = list(dict.fromkeys(source_warnings + warnings))
    report_path = write_report(out, src, "ca", all_warnings)

    content = report_path.read_text(encoding="utf-8")
    assert "Warnings" in content
    assert "Tables detected" in content


def test_conversion_report_clean_message_when_no_warnings(tmp_path: Path) -> None:
    src = tmp_path / "notice.docx"
    out = tmp_path / "notice_ca.docx"
    make_notice(src)

    doc, record = load(src)
    transform(doc)
    save(doc, out)
    report_path = write_report(out, src, "ca", [])

    content = report_path.read_text(encoding="utf-8")
    assert "No warnings" in content
