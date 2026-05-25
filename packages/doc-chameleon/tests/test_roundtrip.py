from __future__ import annotations

from pathlib import Path

import docx as python_docx
import pytest

from doc_chameleon.engine.export import save
from doc_chameleon.engine.ingest import load


def _make_sample_docx(path: Path) -> None:
    doc = python_docx.Document()
    doc.add_heading("Motion for Summary Judgment", level=1)
    doc.add_paragraph("Plaintiff respectfully moves this Court for summary judgment.", style="Normal")
    doc.add_paragraph("There is no genuine dispute as to any material fact.", style="Normal")
    doc.add_heading("Argument", level=2)
    doc.add_paragraph("The standard for summary judgment is well established.", style="Normal")
    doc.save(str(path))


def test_roundtrip_preserves_paragraph_count(tmp_path: Path) -> None:
    source = tmp_path / "source.docx"
    output = tmp_path / "output.docx"
    _make_sample_docx(source)

    doc, source_record = load(source)
    save(doc, output)
    _, output_record = load(output)

    assert len(source_record.paragraphs) == len(output_record.paragraphs)


def test_roundtrip_preserves_paragraph_text(tmp_path: Path) -> None:
    source = tmp_path / "source.docx"
    output = tmp_path / "output.docx"
    _make_sample_docx(source)

    doc, source_record = load(source)
    save(doc, output)
    _, output_record = load(output)

    for src, out in zip(source_record.paragraphs, output_record.paragraphs):
        assert src.text == out.text


def test_roundtrip_preserves_paragraph_styles(tmp_path: Path) -> None:
    source = tmp_path / "source.docx"
    output = tmp_path / "output.docx"
    _make_sample_docx(source)

    doc, source_record = load(source)
    save(doc, output)
    _, output_record = load(output)

    for src, out in zip(source_record.paragraphs, output_record.paragraphs):
        assert src.style_name == out.style_name


def test_roundtrip_preserves_section_count(tmp_path: Path) -> None:
    source = tmp_path / "source.docx"
    output = tmp_path / "output.docx"
    _make_sample_docx(source)

    doc, source_record = load(source)
    save(doc, output)
    _, output_record = load(output)

    assert len(source_record.sections) == len(output_record.sections)


def test_roundtrip_preserves_page_dimensions(tmp_path: Path) -> None:
    source = tmp_path / "source.docx"
    output = tmp_path / "output.docx"
    _make_sample_docx(source)

    doc, source_record = load(source)
    save(doc, output)
    _, output_record = load(output)

    for src, out in zip(source_record.sections, output_record.sections):
        assert src.page_width_emu == out.page_width_emu
        assert src.page_height_emu == out.page_height_emu
        assert src.left_margin_emu == out.left_margin_emu
        assert src.right_margin_emu == out.right_margin_emu
        assert src.top_margin_emu == out.top_margin_emu
        assert src.bottom_margin_emu == out.bottom_margin_emu
