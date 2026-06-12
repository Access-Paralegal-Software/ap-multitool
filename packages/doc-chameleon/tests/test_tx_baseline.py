"""Texas statewide baseline transformer and validator tests."""
from __future__ import annotations

from pathlib import Path

from docx.shared import Inches, Pt

from doc_chameleon.engine.export import save
from doc_chameleon.engine.ingest import load
from doc_chameleon.rules.tx.transformer import TX_MARKER, transform
from doc_chameleon.rules.tx.validator import (
    _STATEWIDE_ONLY_WARNING,
    validate,
    validate_source_assumptions,
)

from tests.fixtures import make_hostile, make_tx_motion, make_tx_notice


def test_tx_transform_applies_page_geometry(tmp_path: Path) -> None:
    src = tmp_path / "source.docx"
    out = tmp_path / "tx.docx"
    make_tx_motion(src)

    doc, _ = load(src)
    transform(doc)
    save(doc, out)

    reloaded, _ = load(out)
    section = reloaded.sections[0]
    assert section.page_width == Inches(8.5)
    assert section.page_height == Inches(11)
    assert section.left_margin == Inches(1)
    assert section.right_margin == Inches(1)
    assert section.top_margin == Inches(1)
    assert section.bottom_margin == Inches(1)


def test_tx_transform_applies_body_style(tmp_path: Path) -> None:
    src = tmp_path / "source.docx"
    out = tmp_path / "tx.docx"
    make_tx_motion(src)

    doc, _ = load(src)
    transform(doc)
    save(doc, out)

    reloaded, _ = load(out)
    for para in reloaded.paragraphs:
        assert para.style.name == "Doc Chameleon TX Body"
        assert para.paragraph_format.space_before == Pt(0)
        assert para.paragraph_format.space_after == Pt(0)


def test_tx_transform_body_font(tmp_path: Path) -> None:
    src = tmp_path / "source.docx"
    make_tx_motion(src)

    doc, _ = load(src)
    transform(doc)

    style = doc.styles["Doc Chameleon TX Body"]
    assert style.font.name == "Times New Roman"
    assert style.font.size == Pt(12)


def test_tx_transform_marks_sectpr(tmp_path: Path) -> None:
    src = tmp_path / "source.docx"
    out = tmp_path / "tx.docx"
    make_tx_motion(src)

    doc, _ = load(src)
    transform(doc)
    save(doc, out)

    reloaded, _ = load(out)
    assert TX_MARKER in reloaded.sections[0]._sectPr.xml


def test_tx_transform_does_not_install_line_number_header(tmp_path: Path) -> None:
    src = tmp_path / "source.docx"
    make_tx_motion(src)

    doc, _ = load(src)
    transform(doc)

    for section in doc.sections:
        assert "doc-chameleon-ca-line-numbers" not in section.header._element.xml


def test_tx_transform_preserves_paragraph_count(tmp_path: Path) -> None:
    src = tmp_path / "source.docx"
    out = tmp_path / "tx.docx"
    make_tx_motion(src)

    doc, record = load(src)
    original_count = len(record.paragraphs)
    transform(doc)
    save(doc, out)

    _, reloaded_record = load(out)
    assert len(reloaded_record.paragraphs) == original_count


def test_tx_validator_always_includes_statewide_only_warning(tmp_path: Path) -> None:
    src = tmp_path / "source.docx"
    make_tx_motion(src)

    doc, record = load(src)
    transform(doc)
    warnings = validate(doc, record)

    assert any(_STATEWIDE_ONLY_WARNING in w for w in warnings)


def test_tx_validator_warns_if_transform_not_applied(tmp_path: Path) -> None:
    src = tmp_path / "source.docx"
    make_tx_motion(src)

    doc, record = load(src)
    warnings = validate(doc, record)

    assert any("transform not applied" in w for w in warnings)


def test_tx_validator_detects_caption_in_tx_motion(tmp_path: Path) -> None:
    src = tmp_path / "source.docx"
    make_tx_motion(src)

    doc, record = load(src)
    transform(doc)
    warnings = validate(doc, record)

    assert not any("caption" in w.lower() for w in warnings)


def test_tx_validator_warns_missing_caption(tmp_path: Path) -> None:
    src = tmp_path / "source.docx"
    doc_raw = __import__("docx").Document()
    doc_raw.add_paragraph("Respondent moves for summary judgment.")
    doc_raw.add_paragraph("The material facts are undisputed.")
    doc_raw.save(str(src))

    doc, record = load(src)
    transform(doc)
    warnings = validate(doc, record)

    assert any("caption" in w.lower() for w in warnings)


def test_tx_notice_fixture_transforms_cleanly(tmp_path: Path) -> None:
    src = tmp_path / "notice.docx"
    out = tmp_path / "notice_tx.docx"
    make_tx_notice(src)

    doc, record = load(src)
    source_warnings = validate_source_assumptions(doc, record)
    assert source_warnings == [], f"Clean notice should have no source warnings: {source_warnings}"

    transform(doc)
    save(doc, out)

    reloaded, _ = load(out)
    assert TX_MARKER in reloaded.sections[0]._sectPr.xml


def test_tx_source_assumptions_flags_explicit_spacing(tmp_path: Path) -> None:
    src = tmp_path / "source.docx"
    import docx as python_docx
    doc_raw = python_docx.Document()
    p = doc_raw.add_paragraph("Spaced paragraph.")
    p.paragraph_format.space_after = Pt(12)
    doc_raw.save(str(src))

    doc, record = load(src)
    warnings = validate_source_assumptions(doc, record)

    assert any("paragraph spacing" in w for w in warnings)
