from __future__ import annotations

from pathlib import Path
from zipfile import ZipFile

import docx as python_docx
from docx.enum.text import WD_LINE_SPACING
from docx.shared import Inches, Pt

from doc_chameleon.engine.export import save
from doc_chameleon.engine.ingest import load
from doc_chameleon.rules.ca.transformer import PLEADING_LINE_COUNT, transform
from doc_chameleon.rules.ca.validator import validate, validate_source_assumptions


def _make_motion_docx(path: Path) -> None:
    doc = python_docx.Document()
    doc.add_paragraph("Motion for Summary Judgment")
    doc.add_paragraph("Plaintiff respectfully moves this Court for summary judgment.")
    doc.add_paragraph("The material facts are undisputed.")
    doc.save(str(path))


def _header_xml(path: Path) -> str:
    with ZipFile(path) as package:
        header_names = [name for name in package.namelist() if name.startswith("word/header")]
        assert header_names
        return "\n".join(package.read(name).decode("utf-8") for name in header_names)


def test_ca_transform_installs_structural_line_number_header(tmp_path: Path) -> None:
    source = tmp_path / "source.docx"
    output = tmp_path / "ca.docx"
    _make_motion_docx(source)

    doc, _record = load(source)
    transform(doc)
    save(doc, output)

    xml = _header_xml(output)
    assert "doc-chameleon-ca-line-numbers" in xml
    assert "<w:txbxContent>" in xml
    assert f"<w:t>{PLEADING_LINE_COUNT}</w:t>" in xml


def test_ca_transform_applies_controlled_page_geometry_and_body_flow(tmp_path: Path) -> None:
    source = tmp_path / "source.docx"
    output = tmp_path / "ca.docx"
    _make_motion_docx(source)

    doc, _record = load(source)
    transform(doc)
    save(doc, output)
    transformed, _transformed_record = load(output)

    section = transformed.sections[0]
    assert section.page_width == Inches(8.5)
    assert section.page_height == Inches(11)
    assert section.left_margin == Inches(1.25)
    assert section.right_margin == Inches(1)
    assert section.top_margin == Inches(1)
    assert section.bottom_margin == Inches(1)

    for paragraph in transformed.paragraphs:
        assert paragraph.style.name == "Doc Chameleon CA Body"
        assert paragraph.paragraph_format.line_spacing_rule == WD_LINE_SPACING.EXACTLY
        assert paragraph.paragraph_format.line_spacing == Pt(24)
        assert paragraph.paragraph_format.space_before == Pt(0)
        assert paragraph.paragraph_format.space_after == Pt(0)


def test_ca_validator_warns_before_numbering_transform(tmp_path: Path) -> None:
    source = tmp_path / "source.docx"
    _make_motion_docx(source)

    doc, record = load(source)
    warnings = validate(doc, record)

    assert any("line-number header not found" in warning for warning in warnings)


def test_ca_validator_flags_tables_for_manual_review(tmp_path: Path) -> None:
    source = tmp_path / "source.docx"
    doc = python_docx.Document()
    doc.add_paragraph("Declaration of Jane Smith")
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "Unsupported layout block"
    doc.save(str(source))

    loaded, record = load(source)
    transform(loaded)
    warnings = validate(loaded, record)

    assert any("Tables detected" in warning for warning in warnings)


def test_ca_source_assumption_validator_flags_manual_spacing_before_transform(tmp_path: Path) -> None:
    source = tmp_path / "source.docx"
    doc = python_docx.Document()
    paragraph = doc.add_paragraph("Manually spaced source paragraph")
    paragraph.paragraph_format.space_after = Pt(12)
    doc.save(str(source))

    loaded, record = load(source)
    warnings = validate_source_assumptions(loaded, record)

    assert any("paragraph spacing" in warning for warning in warnings)
