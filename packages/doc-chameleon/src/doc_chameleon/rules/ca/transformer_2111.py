"""CRC 2.111 first-page layout transformer.

Additive layer intended to run after the 2.108 line-number transform.
Enables a different first-page header and installs the attorney/clerk split.
"""
from __future__ import annotations

from dataclasses import dataclass, field

import docx
from docx.oxml import parse_xml
from docx.shared import Inches
from lxml import etree

from doc_chameleon.rules.ca.transformer import (
    PLEADING_FONT_NAME,
    PLEADING_FONT_SIZE_PT,
    PLEADING_LINE_SPACING_PT,
    _clear_header,
    _line_number_pict,
)

CLERK_SPACE_LABEL = "FOR COURT USE ONLY"

_CA_2111_MARKER = "doc-chameleon-ca-2111"
_ATTORNEY_COL_TWIPS = int(Inches(3.5).pt * 20)   # 5040 twips = 3.5"
_CLERK_COL_TWIPS = int(Inches(2.75).pt * 20)      # 3960 twips = 2.75"


@dataclass
class AttorneyInfo:
    """Attorney information for the CRC 2.111 first-page header block.

    All fields default to bracket placeholders. Provide specific values to
    pre-fill the attorney block instead of requiring manual post-conversion edits.
    """
    name: str = "[Attorney Name]"
    sbn: str = "[XXXXX]"
    firm: str = "[Firm Name]"
    address: str = "[Street Address]"
    city_state_zip: str = "[City, State ZIP]"
    phone: str = "[Phone]"
    email: str = "[email@domain.com]"
    client: str = "[Plaintiff/Defendant], [Party Name]"


def _attorney_lines(attorney: AttorneyInfo) -> list[str]:
    return [
        f"{attorney.name}, SBN {attorney.sbn}",
        attorney.firm,
        attorney.address,
        attorney.city_state_zip,
        f"Tel: {attorney.phone}",
        f"Email: {attorney.email}",
        "",
        f"Attorney for {attorney.client}",
    ]


# Module-level constant for test assertions — always matches AttorneyInfo defaults.
ATTORNEY_PLACEHOLDER_LINES: list[str] = _attorney_lines(AttorneyInfo())


def transform_2111(doc: docx.Document, attorney: AttorneyInfo | None = None) -> docx.Document:
    if attorney is None:
        attorney = AttorneyInfo()
    section = doc.sections[0]
    section.different_first_page_header_footer = True
    _install_first_page_header(doc, attorney)
    return doc


def _install_first_page_header(doc: docx.Document, attorney: AttorneyInfo) -> None:
    section = doc.sections[0]
    header = section.first_page_header
    _clear_header(header)

    ln_para = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
    ln_para._p.append(_line_number_pict())

    header._element.append(_attorney_clerk_table(_attorney_lines(attorney)))
    header._element.append(etree.Comment(_CA_2111_MARKER))


def _attorney_clerk_table(lines: list[str]):
    rows_xml = ""
    for i, line in enumerate(lines):
        attorney_text = line.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
        clerk_text = CLERK_SPACE_LABEL if i == 0 else ""
        line_twips = PLEADING_LINE_SPACING_PT * 20
        font_half_pts = PLEADING_FONT_SIZE_PT * 2

        rows_xml += f"""
        <w:tr>
          <w:tc>
            <w:tcPr>
              <w:tcW w:w="{_ATTORNEY_COL_TWIPS}" w:type="dxa"/>
            </w:tcPr>
            <w:p>
              <w:pPr>
                <w:spacing w:before="0" w:after="0" w:line="{line_twips}" w:lineRule="exact"/>
              </w:pPr>
              <w:r>
                <w:rPr>
                  <w:rFonts w:ascii="{PLEADING_FONT_NAME}" w:hAnsi="{PLEADING_FONT_NAME}"/>
                  <w:sz w:val="{font_half_pts}"/>
                </w:rPr>
                <w:t xml:space="preserve">{attorney_text}</w:t>
              </w:r>
            </w:p>
          </w:tc>
          <w:tc>
            <w:tcPr>
              <w:tcW w:w="{_CLERK_COL_TWIPS}" w:type="dxa"/>
            </w:tcPr>
            <w:p>
              <w:pPr>
                <w:spacing w:before="0" w:after="0" w:line="{line_twips}" w:lineRule="exact"/>
              </w:pPr>
              <w:r>
                <w:rPr>
                  <w:rFonts w:ascii="{PLEADING_FONT_NAME}" w:hAnsi="{PLEADING_FONT_NAME}"/>
                  <w:sz w:val="{font_half_pts}"/>
                </w:rPr>
                <w:t xml:space="preserve">{clerk_text}</w:t>
              </w:r>
            </w:p>
          </w:tc>
        </w:tr>"""

    total_twips = _ATTORNEY_COL_TWIPS + _CLERK_COL_TWIPS

    return parse_xml(
        f'<w:tbl xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        f'<w:tblPr>'
        f'  <w:tblW w:w="{total_twips}" w:type="dxa"/>'
        f'  <w:tblBorders>'
        f'    <w:top w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
        f'    <w:left w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
        f'    <w:bottom w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
        f'    <w:right w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
        f'    <w:insideH w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
        f'    <w:insideV w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
        f'  </w:tblBorders>'
        f'</w:tblPr>'
        f'<w:tblGrid>'
        f'  <w:gridCol w:w="{_ATTORNEY_COL_TWIPS}"/>'
        f'  <w:gridCol w:w="{_CLERK_COL_TWIPS}"/>'
        f'</w:tblGrid>'
        f'{rows_xml}'
        f'</w:tbl>'
    )
