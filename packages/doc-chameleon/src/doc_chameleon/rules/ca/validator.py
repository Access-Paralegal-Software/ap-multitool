from __future__ import annotations

import docx
from docx.enum.text import WD_LINE_SPACING
from docx.shared import Inches, Pt

from doc_chameleon.engine.models import DocumentRecord

EXPECTED_PAGE_WIDTH = Inches(8.5)
EXPECTED_PAGE_HEIGHT = Inches(11)
EXPECTED_LEFT_MARGIN = Inches(1.25)
EXPECTED_RIGHT_MARGIN = Inches(1)
EXPECTED_TOP_MARGIN = Inches(1)
EXPECTED_BOTTOM_MARGIN = Inches(1)


def validate(doc: docx.Document, record: DocumentRecord) -> list[str]:
    """Validate document against the controlled California numbering assumptions."""
    warnings: list[str] = []

    for index, section in enumerate(doc.sections, start=1):
        warnings.extend(_validate_section(index, section))

    if len(doc.sections) > 1:
        warnings.append(
            "Multiple sections detected. California numbering spike applies one controlled page geometry to every section; verify section breaks manually."
        )

    warnings.extend(validate_source_assumptions(doc, record))
    if not _has_ca_line_number_header(doc):
        warnings.append("California line-number header not found. Run the CA transformer before relying on CRC 2.108 numbering.")

    return warnings


def validate_source_assumptions(doc: docx.Document, record: DocumentRecord) -> list[str]:
    """Warn on source structures that do not cleanly fit the controlled CA body flow."""
    warnings: list[str] = []

    if doc.tables:
        warnings.append("Tables detected. Table-heavy layouts can break visual line-number alignment and require manual review.")

    if _has_drawings_or_embeds(doc):
        warnings.append("Images, drawings, text boxes, or embedded objects detected. Layout-affecting objects require manual review.")

    for paragraph in doc.paragraphs:
        fmt = paragraph.paragraph_format
        if fmt.line_spacing_rule not in (None, WD_LINE_SPACING.EXACTLY):
            warnings.append(
                f"Paragraph {paragraph_index(record, paragraph.text)} uses non-exact line spacing; California line alignment assumes exact 24 pt spacing."
            )
            break
        if fmt.space_before not in (None, Pt(0)) or fmt.space_after not in (None, Pt(0)):
            warnings.append(
                f"Paragraph {paragraph_index(record, paragraph.text)} has paragraph spacing before/after; California line alignment assumes zero paragraph spacing."
            )
            break

    return warnings


def _validate_section(index: int, section: docx.section.Section) -> list[str]:
    warnings: list[str] = []
    checks = [
        ("page width", section.page_width, EXPECTED_PAGE_WIDTH),
        ("page height", section.page_height, EXPECTED_PAGE_HEIGHT),
        ("left margin", section.left_margin, EXPECTED_LEFT_MARGIN),
        ("right margin", section.right_margin, EXPECTED_RIGHT_MARGIN),
        ("top margin", section.top_margin, EXPECTED_TOP_MARGIN),
        ("bottom margin", section.bottom_margin, EXPECTED_BOTTOM_MARGIN),
    ]

    for label, actual, expected in checks:
        if actual != expected:
            warnings.append(
                f"Section {index} {label} is {actual} EMU; controlled California numbering expects {expected} EMU."
            )

    return warnings


def _has_drawings_or_embeds(doc: docx.Document) -> bool:
    body_xml = doc._element.body.xml
    return any(marker in body_xml for marker in ("<w:drawing", "<w:pict", "<w:object", "<w:txbxContent"))


def _has_ca_line_number_header(doc: docx.Document) -> bool:
    return any("doc-chameleon-ca-line-numbers" in section.header._element.xml for section in doc.sections)


def paragraph_index(record: DocumentRecord, text: str) -> int | str:
    for paragraph in record.paragraphs:
        if paragraph.text == text:
            return paragraph.index
    return "unknown"
