"""Word → PDF conversion with COM primary and python-docx fallback.

When Microsoft Word is not installed (or win32com / COM fails), conversion
falls back to extracting text via python-docx and rendering with ReportLab.
This path is intentionally layout-limited but keeps headless/air-gapped
machines functional without Office.
"""

from __future__ import annotations

from pathlib import Path
from xml.sax.saxutils import escape as xml_escape

from ap_core.logging_config import get_logger, safe_filename

logger = get_logger("core.converters.docx")

# COM / Word-not-installed failures we treat as fallback triggers.
_COM_FALLBACK_EXCEPTIONS: tuple[type[BaseException], ...] = (ImportError, OSError)


def _com_error_types() -> tuple[type[BaseException], ...]:
    """Return pywintypes.com_error when available (Windows + pywin32)."""
    try:
        import pywintypes  # type: ignore

        return (pywintypes.com_error,)
    except ImportError:
        return ()


def _convert_via_win32com(src_path: Path, out_path: Path) -> None:
    """Export via Microsoft Word COM automation. Raises on any failure."""
    import pythoncom
    import win32com.client

    pythoncom.CoInitialize()
    word = None
    doc = None
    try:
        word = win32com.client.DispatchEx("Word.Application")
        word.Visible = False
        word.DisplayAlerts = False
        doc = word.Documents.Open(str(src_path.resolve()), ReadOnly=True)
        doc.SaveAs(str(out_path.resolve()), FileFormat=17)  # 17 = wdFormatPDF
    finally:
        if doc:
            try:
                doc.Close(SaveChanges=0)
            except Exception:
                pass
        if word:
            try:
                word.Quit()
            except Exception:
                pass
        try:
            pythoncom.CoUninitialize()
        except Exception:
            pass


def _convert_via_python_docx(src_path: Path, out_path: Path) -> None:
    """Render .docx text content to PDF via python-docx + ReportLab.

    .doc (legacy binary) is not supported by python-docx and will raise.
    """
    import docx as _docx
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import inch
    from reportlab.platypus import (
        HRFlowable,
        Paragraph,
        SimpleDocTemplate,
        Spacer,
    )

    if src_path.suffix.lower() == ".doc":
        raise RuntimeError(
            "python-docx cannot open legacy .doc files; "
            "install Microsoft Word or LibreOffice for .doc conversion."
        )

    styles = getSampleStyleSheet()

    def _style(name: str, **kwargs) -> ParagraphStyle:
        return ParagraphStyle(name, parent=styles["Normal"], **kwargs)

    body = _style("Body", fontName="Helvetica", fontSize=10, leading=14, spaceAfter=4)
    h1 = _style("H1", fontName="Helvetica-Bold", fontSize=13, leading=18, spaceAfter=6)
    h2 = _style("H2", fontName="Helvetica-Bold", fontSize=11, leading=15, spaceAfter=4)
    label = _style("Label", fontName="Helvetica-Bold", fontSize=11, spaceAfter=8)

    doc_in = _docx.Document(str(src_path))
    story: list = [
        Paragraph(xml_escape(src_path.name), label),
        HRFlowable(width="100%", thickness=0.5, color=colors.grey, spaceAfter=6),
    ]
    for para in doc_in.paragraphs:
        txt = para.text
        if not txt.strip():
            story.append(Spacer(1, 0.08 * inch))
            continue
        style_name = para.style.name or ""
        if "Heading 1" in style_name:
            st = h1
        elif "Heading 2" in style_name:
            st = h2
        else:
            st = body
        story.append(Paragraph(xml_escape(txt), st))

    # Tables: flatten cell text so tabular content is not silently dropped.
    for table in doc_in.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            line = " | ".join(c for c in cells if c)
            if line:
                story.append(Paragraph(xml_escape(line), body))

    pdf = SimpleDocTemplate(
        str(out_path),
        pagesize=letter,
        leftMargin=inch,
        rightMargin=inch,
        topMargin=inch,
        bottomMargin=inch,
    )
    pdf.build(story)


def convert_docx_file_to_pdf(src_path: Path, out_path: Path) -> list[str]:
    """Convert a Word document to PDF with COM → python-docx failover.

    Returns a list of warning strings (non-empty when the python-docx path was used).

    Raises:
        FileNotFoundError: If *src_path* does not exist.
        RuntimeError: If both COM and python-docx paths fail.
    """
    if not src_path.exists():
        raise FileNotFoundError(f"Source file not found: {src_path}")

    warnings: list[str] = []
    fallback_exc: BaseException | None = None
    catch = _COM_FALLBACK_EXCEPTIONS + _com_error_types() + (Exception,)

    try:
        _convert_via_win32com(src_path, out_path)
        logger.info(
            "docx_converter_completed backend=win32com input_filename=%s",
            safe_filename(src_path),
        )
        return warnings
    except catch as com_exc:
        fallback_exc = com_exc
        logger.warning(
            "docx_converter_com_failed input_filename=%s reason=%s fallback=python-docx",
            safe_filename(src_path),
            type(com_exc).__name__,
        )

    try:
        _convert_via_python_docx(src_path, out_path)
        warnings.append(
            f"Microsoft Word COM unavailable ({fallback_exc}). "
            "Used python-docx text extraction fallback."
        )
        logger.info(
            "docx_converter_completed backend=python-docx input_filename=%s",
            safe_filename(src_path),
        )
        return warnings
    except Exception as docx_exc:
        logger.error(
            "docx_converter_failed backends=win32com,python-docx input_filename=%s "
            "com_reason=%s docx_reason=%s",
            safe_filename(src_path),
            type(fallback_exc).__name__ if fallback_exc else "none",
            type(docx_exc).__name__,
        )
        raise RuntimeError(
            f"Word conversion failed. COM: {fallback_exc}; python-docx: {docx_exc}"
        ) from docx_exc
